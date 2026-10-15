"""Offline text rendition / archive extraction. Never execute macros or fetch links.
Outputs derivatives, never modifies originals. Node owns CMS verification and PDF OCR.
"""
import sys, os, json, zipfile, pathlib, textwrap, resource, subprocess
import xml.etree.ElementTree as ET
resource.setrlimit(resource.RLIMIT_AS, (768*1024**2, 768*1024**2))
cpu_hard=resource.getrlimit(resource.RLIMIT_CPU)[1]
resource.setrlimit(resource.RLIMIT_CPU, (min(100,cpu_hard) if cpu_hard>0 else 100, min(110,cpu_hard) if cpu_hard>0 else 110))
src, out, mode = sys.argv[1:4]
if mode == 'archive':
    import libarchive
    entries=[]; total=0; budget=int(sys.argv[4]) if len(sys.argv)>4 else 2*1024**3
    with libarchive.file_reader(src) as archive:
        for i, entry in enumerate(archive):
            name=str(entry.pathname)
            if entry.isdir: continue
            if len(entries)>=1000: raise ValueError('Arhiva depășește 1000 de fișiere; necesită procesare separată.')
            # No names supplied by an archive are used as filesystem paths.
            if not entry.isfile or '..' in pathlib.PurePosixPath(name.replace('\\','/')).parts or name.startswith(('/', '\\')):
                entries.append({'name':name,'error':'Intrare specială sau cale nesigură; nu a fost extrasă.'}); continue
            target=os.path.join(out,str(i))
            with open(target,'xb') as f:
                for block in entry.get_blocks():
                    total+=len(block)
                    if total>budget: raise ValueError('Arhiva decomprimată depășește bugetul de 2 GiB; necesită procesare separată.')
                    f.write(block)
            entries.append({'name':name,'path':target,'size':os.path.getsize(target)})
    print(json.dumps(entries)); sys.exit(0)

def lines():
    if mode=='docx':
        with zipfile.ZipFile(src) as z:
            # Includes tables, headers and footnotes; no relationships are followed.
            parts=[p for p in z.namelist() if p=='word/document.xml' or p.startswith(('word/header','word/footer','word/footnotes','word/endnotes')) and p.endswith('.xml')]
            for part in parts:
                if z.getinfo(part).file_size>128*1024**2: raise ValueError('Conținut XML prea mare pentru conversie.')
                yield '['+part+']'
                with z.open(part) as f:
                    for _, e in ET.iterparse(f,events=('end',)):
                        if e.tag.endswith('}p'):
                            yield ''.join((x.text or '') if x.tag.endswith('}t') else '\t' if x.tag.endswith('}tab') else '\n' if x.tag.endswith('}br') else '' for x in e.iter())
                            e.clear()
    elif mode=='xlsx':
        import openpyxl
        with zipfile.ZipFile(src) as z:
            if sum(x.file_size for x in z.infolist())>512*1024**2: raise ValueError('Registrul decomprimat necesită procesare separată.')
        book=openpyxl.load_workbook(open(src,"rb"),read_only=True,data_only=False,keep_links=False)
        cached=openpyxl.load_workbook(open(src,"rb"),read_only=True,data_only=True,keep_links=False)
        for sheet in book:
            yield 'FOAIA: '+sheet.title
            for index,(row,values) in enumerate(zip(sheet.iter_rows(values_only=True),cached[sheet.title].iter_rows(values_only=True)),1):
                cells=[]
                for column,x in enumerate(row):
                    if x is None: continue
                    value=str(x)
                    if isinstance(x,str) and x.startswith('='): value+=' [valoare memorată: '+(str(values[column]) if column<len(values) and values[column] is not None else 'indisponibilă')+']'
                    cells.append(openpyxl.utils.get_column_letter(column+1)+str(index)+' = '+value)
                if cells: yield ' | '.join(cells)
        book.close();cached.close()
    elif mode=='xls':
        import xlrd
        book=xlrd.open_workbook(src,on_demand=True)
        for sheet in book.sheets():
            yield 'FOAIA: '+sheet.name
            for i in range(sheet.nrows): yield str(i+1)+': '+' | '.join(str(x) for x in sheet.row_values(i))
        book.release_resources()
    elif mode=='doc':
        temporary=out+'.txt'
        with open(temporary,'wb') as f: subprocess.run(['antiword','-m','UTF-8.txt',src],stdout=f,stderr=subprocess.DEVNULL,check=True,timeout=90)
        with open(temporary,encoding='utf-8',errors='replace') as f: yield from f
        os.unlink(temporary)
    else: raise ValueError('Format fără convertor.')
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
pdfmetrics.registerFont(TTFont('SourceText','/usr/share/fonts/truetype/dejavu/DejaVuSansMono.ttf'))
pdf=canvas.Canvas(out,pagesize=(595,842),pageCompression=1); page=0; y=0
def newpage():
    global page,y
    if page: pdf.showPage()
    page+=1; y=788
    pdf.setFont('SourceText',8);pdf.drawString(30,820,'Text rearanjat pentru lectură · verifică formatarea în original')
    pdf.drawString(30,18,str(page));pdf.setFont('SourceText',9)
newpage()
for line in lines():
    for segment in str(line).replace('\x00','').splitlines() or ['']:
        for part in textwrap.wrap(segment,96,replace_whitespace=False,drop_whitespace=False) or ['']:
            if y<40: newpage()
            pdf.drawString(30,y,part); y-=12
pdf.save()
print(json.dumps({'pages':page,'rendition':'text'}))
